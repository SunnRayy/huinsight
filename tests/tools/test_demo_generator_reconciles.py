"""The demo's history and its latest point must describe the same portfolio.

The dashboard trend line is the Financial Summary's 合计总资产 and its last
point is the sum of reader holdings. If the generator's FS fund columns do not
end where the reader files do, a fresh demo shows a phantom jump or crash at
the last point (Round 5 #2: −42% with prices down, +14% with them up).
"""
import importlib.util
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(scope="module")
def generated(tmp_path_factory):
    spec = importlib.util.spec_from_file_location("demo_generate", ROOT / "tools" / "demo_data" / "generate.py")
    gen = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(gen)
    out = tmp_path_factory.mktemp("demo_out")
    gen.generate_all(gen.load_persona(gen.PERSONA_PATH), out)
    return gen, out


def _last_bs_value(fs_path: Path, column: str) -> float:
    """Last non-empty value of a column on the 资产负债 sheet (header row located by name)."""
    ws = openpyxl.load_workbook(fs_path, data_only=True)["资产负债"]
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == column:
                values = [
                    ws.cell(row=r, column=cell.column).value
                    for r in range(cell.row + 1, ws.max_row + 1)
                ]
                return float([v for v in values if v is not None][-1])
    raise AssertionError(f"column {column} not found in {fs_path.name}")


def _cn_fund_holdings_value(path: Path) -> float:
    """Sum of the latest 基金持仓汇总 rows: [code, name, type, date, nav, qty, value]."""
    rows = list(openpyxl.load_workbook(path, data_only=True)["基金持仓汇总"].iter_rows(min_row=2, values_only=True))
    latest = max(r[3] for r in rows if r[3] is not None)
    return sum(float(r[6]) for r in rows if r[3] == latest)


def test_fs_cn_fund_column_ends_at_the_cn_fund_holdings(generated):
    _gen, out = generated
    fs_value = _last_bs_value(out / "Financial_Summary_new.xlsx", "投资资产_股票基金_A股基金")
    held = _cn_fund_holdings_value(out / "funding_transactions.xlsx")
    assert fs_value == pytest.approx(held, rel=1e-6)


def test_cn_fund_navs_sit_at_their_persona_anchors(generated):
    gen, out = generated
    persona = gen.load_persona(gen.PERSONA_PATH)
    anchors = {f["code"]: f.get("nav", 1.0 if f["type"] == "货币型" else None) for f in persona["cn_funds"]["catalog"]}
    ws = openpyxl.load_workbook(out / "funding_transactions.xlsx", data_only=True)["基金持仓汇总"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    snapshot = max(r[3] for r in rows)
    for code, name, _type, d, price, *_rest in rows:
        if d == snapshot and anchors.get(code) is not None:
            assert float(price) == pytest.approx(anchors[code]), code
